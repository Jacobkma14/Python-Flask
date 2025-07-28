import openpyxl

wb = openpyxl.Workbook()
sheet = wb.active

weather_data = [["Day", "Temperature (°C)", "Humidity (%)", "Wind Speed (km/h)"],
                ["Monday", 22, 60, 15],
                ["Tuesday", 24, 55, 18], 
                ["Wednesday", 26, 50, 20],
                ["Thursday", 23, 65, 14], 
                ["Friday", 21, 70, 12],
                ["Saturday", 25, 45, 22], 
                ["Sunday", 27, 40, 25]]

for row in weather_data:  # Add data to the sheet
    sheet.append(row)

# For Temperature Bar Chart
refObj_temp = openpyxl.chart.Reference(sheet, min_col=2, min_row=1, max_row=8)
seriesObj_temp = openpyxl.chart.Series(refObj_temp, title="Temperature")
chartObj_temp = openpyxl.chart.BarChart()
chartObj_temp.title = "Temperature Over the Week"
chartObj_temp.append(seriesObj_temp)
sheet.add_chart(chartObj_temp, "G2")

# For Humidity Line Chart
refObj_humidity = openpyxl.chart.Reference(sheet,min_col=3,min_row=1,max_row=8)
seriesObj_humidity = openpyxl.chart.Series(refObj_humidity, title="Humidity")
chartObj_humidity = openpyxl.chart.LineChart()
chartObj_humidity.title = "Humidity Over the Week"
chartObj_humidity.append(seriesObj_humidity)
sheet.add_chart(chartObj_humidity, "G20")

# For Temperature vs Wind Speed Scatter Chart
chartObj_temp_wind = openpyxl.chart.ScatterChart()
chartObj_temp_wind.title = "Temperature vs Wind Speed"
chartObj_temp_wind.x_axis.title = "Temperature (°C)"
chartObj_temp_wind.y_axis.title = "Wind Speed (km/h)"
x_values = openpyxl.chart.Reference(sheet, min_col=2, min_row=2, max_row=8)  # Temperature (X-axis)
y_values = openpyxl.chart.Reference(sheet, min_col=4, min_row=2, max_row=8)  # Wind Speed (Y-axis)
series = openpyxl.chart.Series(y_values, xvalues=x_values, title="Wind Speed vs Temperature")
chartObj_temp_wind.series.append(series)
series.marker.symbol = "circle" 
sheet.add_chart(chartObj_temp_wind, "Q2")

# For Humidity Pie Chart
refObj_pie = openpyxl.chart.Reference(sheet, min_col=3, min_row=2, max_row=8)
seriesObj_pie = openpyxl.chart.Series(refObj_pie, title="Humidity")
chartObj_pie = openpyxl.chart.PieChart()
chartObj_pie.title = "Humidity Distribution"
chartObj_pie.append(seriesObj_pie)
sheet.add_chart(chartObj_pie, "Q20")

file_name = "Weather_Report.xlsx"
wb.save(file_name)
print(f"Excel file '{file_name}' has been created successfully!")
